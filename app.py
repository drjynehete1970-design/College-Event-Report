"""Run with: streamlit run app.py"""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from xml.sax.saxutils import escape

import streamlit as st
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from PIL import Image as PILImage, ImageOps
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, KeepTogether, Table, TableStyle
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4


FIELDS = [
    'Event Title', 'Date & Time', 'Venue', 'Agenda', 'Organized by', 'Coordinator',
    'Resource Person/Guest', 'Participants', 'Objectives', 'Brief Report', 'Outcome',
]


def normalize_image(raw):
    with PILImage.open(BytesIO(raw)) as source:
        im = ImageOps.exif_transpose(source).convert('RGB')
        im.thumbnail((1800, 1800))
        output = BytesIO()
        im.save(output, format='JPEG')
        return output.getvalue(), im.width, im.height


def export_report(college, values, photos, attachments, signatures, activity_reference=''):
    if len(photos) != 3 or any(not caption.strip() for _, caption in photos):
        raise ValueError('Supply exactly three photographs, each with a caption.')
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.8)
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    reference_line = f'Activity Reference Number: {activity_reference.strip()}'
    doc.add_paragraph(reference_line)
    doc.add_heading(college, 0)
    pdf = BytesIO()
    styles = getSampleStyleSheet()
    story = [Paragraph(escape(reference_line), styles['BodyText']), Spacer(1, 8),
             Paragraph(escape(college), styles['Title'])]

    def paragraph(text, style='BodyText'):
        return Paragraph(escape(text).replace('\n', '<br/>'), styles[style])

    def add_section(number, title, body):
        heading = f'{number}. {title}'
        doc.add_heading(heading, 1)
        doc.add_paragraph(body or 'Not provided')
        story.extend([paragraph(heading, 'Heading2'), paragraph(body or 'Not provided'), Spacer(1, 8)])

    for i, label in enumerate(FIELDS, 1):
        add_section(i, label, values[label])
    doc.add_heading('12. Photographs', 1)
    story.append(paragraph('12. Photographs', 'Heading2'))
    for raw, caption in photos:
        data, width, height = normalize_image(raw)
        scale = min(430 / width, 245 / height)
        doc.add_picture(BytesIO(data), width=Inches(width * scale / 72), height=Inches(height * scale / 72))
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
        doc.add_paragraph(caption, 'Caption')
        story.append(KeepTogether([Image(BytesIO(data), width=width * scale, height=height * scale),
                                  paragraph(caption), Spacer(1, 12)]))
    names = '\n'.join(name for name, _ in attachments)
    add_section(13, 'Supporting Documents', names or 'No supporting documents supplied.')
    if attachments:
        note = 'Original supporting documents are included in the report package ZIP.'
        doc.add_paragraph(note)
        story.append(paragraph(note))
    if signatures:
        # A single borderless row keeps each name aligned with its role.
        table = doc.add_table(rows=1, cols=len(signatures))
        table.autofit = False
        width = (section.page_width - section.left_margin - section.right_margin) // len(signatures)
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for column, cell, (role, name) in zip(table.columns, table.rows[0].cells, signatures):
            column.width = width
            cell.width = width
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(36)
            p.paragraph_format.space_after = Pt(0)
            p.add_run(name.strip())
            p.add_run('\n' + role)
        signature_style = ParagraphStyle('Signature', parent=styles['BodyText'], alignment=1)
        cells = [Paragraph(escape(name.strip()) + '<br/>' + escape(role), signature_style)
                 for role, name in signatures]
        pdf_table = Table([cells], colWidths=[(A4[0] - 100) / len(signatures)] * len(signatures))
        pdf_table.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'),
                                      ('TOPPADDING', (0, 0), (-1, -1), 36),
                                      ('BOTTOMPADDING', (0, 0), (-1, -1), 0)]))
        story.append(pdf_table)
    else:
        doc.add_paragraph('Not provided')
        story.append(paragraph('Not provided'))
    word = BytesIO()
    doc.save(word)

    def footer(canvas, document):
        canvas.setFont('Helvetica', 9)
        canvas.drawRightString(A4[0] - 50, 25, f'Page {document.page}')

    SimpleDocTemplate(pdf, pagesize=A4, rightMargin=50, leftMargin=50,
                      topMargin=45, bottomMargin=45).build(story, onFirstPage=footer, onLaterPages=footer)
    package = BytesIO()
    with ZipFile(package, 'w') as archive:
        archive.writestr('event_report.docx', word.getvalue())
        archive.writestr('event_report.pdf', pdf.getvalue())
        for i, (name, data) in enumerate(attachments, 1):
            archive.writestr(f'supporting_documents/{i:02d}_{Path(name).name}', data)
    return word.getvalue(), pdf.getvalue(), package.getvalue()


def main():
    st.set_page_config(page_title='College Event Reports', page_icon='🎓')
    st.title('College Event Report Generator')
    st.write('Complete your college’s 14-section format, then download Word, PDF, or a full report package.')
    st.caption('Enter factual event details. Reports reproduce your text; they do not invent activities or outcomes. PDF export currently supports English text.')
    activity_reference = st.text_input('Activity Reference Number:')
    college = st.text_input('College name')
    values = {}
    hints = {'Agenda': 'List the planned activities and their timings, one per line',
             'Date & Time': 'Date, start/end times and duration',
             'Resource Person/Guest': 'Name, designation and organization; enter Not applicable if none',
             'Participants': 'Students, faculty, others and total count',
             'Objectives': 'Enter 2–4 specific objectives, one per line',
             'Brief Report': 'Describe activities, key points and highlights',
             'Outcome': 'Describe what participants gained'}
    for i, label in enumerate(FIELDS, 1):
        values[label] = st.text_area(f'{i}. {label}', placeholder=hints.get(label, ''), height=100) if label in {'Agenda', 'Objectives', 'Brief Report', 'Outcome'} else st.text_input(f'{i}. {label}', placeholder=hints.get(label, ''))
    st.subheader('12. Photographs')
    st.caption('Add exactly three photographs and a caption for each.')
    photos = []
    for i in range(1, 4):
        file = st.file_uploader(f'Photograph {i}', type=['jpg', 'jpeg', 'png'], key=f'photo_{i}')
        caption = st.text_input(f'Caption for photograph {i}', key=f'photo_caption_{i}')
        if file is not None:
            photos.append((file.getvalue(), caption))
    st.subheader('13. Supporting Documents')
    files = st.file_uploader('Notice, attendance, invitation, feedback or certificates',
                             type=['pdf', 'docx', 'xlsx', 'csv', 'png', 'jpg', 'jpeg'], accept_multiple_files=True)
    st.caption('Files are listed in the report and bundled in the ZIP; their contents are not merged into the report.')
    st.subheader('14. Signatures')
    signatures = []
    for role in ['Coordinator', 'IQAC Coordinator', 'Principal']:
        if st.checkbox(f'Include {role}', value=True):
            signatures.append((role, st.text_input(f'{role} name')))
    if st.button('Generate report', type='primary'):
        st.session_state.pop('generated', None)
        if not college.strip() or any(not value.strip() for value in values.values()):
            st.error('Enter the college name and sections 1–11. Use Not applicable where appropriate.')
        elif len(photos) != 3:
            st.error('Please supply all three photographs.')
        elif any(not caption.strip() for _, caption in photos):
            st.error('Add a caption for every photograph.')
        else:
            try:
                st.session_state.generated = export_report(college, values, photos,
                    [(file.name, file.getvalue()) for file in files], signatures, activity_reference)
            except Exception as error:
                st.error(f'Report generation failed. Check your uploaded images and text. Details: {error}')
    if 'generated' in st.session_state:
        st.success('Report generated. After changing inputs, click Generate report again to update downloads.')
        word, pdf, package = st.session_state.generated
        st.download_button('Download Word', word, 'event_report.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        st.download_button('Download PDF', pdf, 'event_report.pdf', 'application/pdf')
        st.download_button('Download full package', package, 'event_report_package.zip', 'application/zip')


if __name__ == '__main__':
    main()
